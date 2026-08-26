from pathlib import Path

from openpyxl import load_workbook

from ttn_bot.excel import DEFAULT_EXCEL_TEMPLATE_PATH, write_ttn_excel
from ttn_bot.models import TtnEditData
from ttn_bot.pdf.extractor import ExtractedItem, ExtractedTtnData


def test_write_ttn_excel_creates_structured_workbook(tmp_path: Path) -> None:
    output = tmp_path / "ttn.xlsx"
    edit_data = TtnEditData(
        vehicle="Газель NN ВА9712-7",
        driver="Мартиневский Иван Васильевич",
        ttn_series="ЕМ",
        ttn_number="1926709",
        goods_accepted_by="Мартиневский Иван Васильевич",
        release_allowed_by="Китасов А.С.",
        shipper_handed_over_by="Китасов А.С.",
    )
    extracted = ExtractedTtnData(
        text="ТОВАРНЫЙ РАЗДЕЛ",
        items=[ExtractedItem(row_number=1, name="Товар 1", quantity="2", unit="шт")],
        tables=[
            [
                ["Наименование товара", "Единица измерения", "Количество"],
                ["Товар 1", "шт.", "2"],
            ]
        ],
    )

    write_ttn_excel(output, edit_data, extracted, template_path=None)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Table 1"]
    assert workbook["Table 1"]["A1"].value == "Наименование товара"
    assert workbook["Table 1"]["A2"].value == "Товар 1"
    assert workbook["Table 1"]["C2"].value == 2


def test_write_ttn_excel_preserves_sample_template_format(tmp_path: Path) -> None:
    output = tmp_path / "ttn_from_template.xlsx"
    edit_data = TtnEditData(
        vehicle="Газель Next ВА5586-7",
        driver="Лабун Александр Чеславович",
        ttn_series="ЕМ",
        ttn_number="1926709",
        goods_accepted_by="Лабун Александр Чеславович",
        release_allowed_by="Шараев А.А.",
        shipper_handed_over_by="Шараев А.А.",
    )
    extracted = ExtractedTtnData(
        text=(
            "02 августа 2026 г.\n"
            "ТОВАРНЫЙ РАЗДЕЛ\n"
            "Всего сумма НДС ноль руб. 0 коп.\n"
            "Всего стоимость c НДС четыре тысячи сто сорок восемь руб. 95 коп.\n"
            "Всего масса груза шесть кг 621 г\n"
            "Всего количество грузовых мест восемьдесят три"
        ),
        items=[],
        tables=[
            [
                ["Грузоотправитель", "Грузополучатель", "Заказчик"],
                ["193018727", "193602362", "193018727"],
            ],
            [["Наименование товара", "Единица измерения", "Количество"], ["Товар", "шт.", "1"]],
            [["Операции", "Исполнитель"], ["Погрузка", "ООО"]],
            [["Расстояние", "Код"], ["", ""]],
            [["Наименование товара", "Единица измерения", "Количество"], ["Товар 1", "шт.", "2"]],
        ],
    )

    write_ttn_excel(output, edit_data, extracted)

    expected = load_workbook(DEFAULT_EXCEL_TEMPLATE_PATH)
    actual = load_workbook(output)
    assert actual.sheetnames == expected.sheetnames
    assert actual["Table 3"]["A1"].value.startswith("02 августа 2026 г.\nАвтомобиль Газель Next")
    assert actual["Table 11"]["A1"].value is not None
    assert "Шараев А.А." in actual["Table 11"]["A1"].value
    assert "Лабун Александр Чеславович" in actual["Table 11"]["B1"].value

    for expected_sheet in expected.worksheets:
        actual_sheet = actual[expected_sheet.title]
        assert actual_sheet.max_row == expected_sheet.max_row
        assert actual_sheet.max_column == expected_sheet.max_column
        assert list(actual_sheet.merged_cells.ranges) == list(expected_sheet.merged_cells.ranges)
        for column_name, expected_dimension in expected_sheet.column_dimensions.items():
            assert actual_sheet.column_dimensions[column_name].width == expected_dimension.width
        for row_index, expected_dimension in expected_sheet.row_dimensions.items():
            assert actual_sheet.row_dimensions[row_index].height == expected_dimension.height
        for row in range(1, expected_sheet.max_row + 1):
            for column in range(1, expected_sheet.max_column + 1):
                assert (
                    actual_sheet.cell(row, column)._style
                    == expected_sheet.cell(row, column)._style
                )


def test_write_ttn_excel_extends_product_row_formatting(tmp_path: Path) -> None:
    output = tmp_path / "ttn_long_product_table.xlsx"
    edit_data = TtnEditData(
        vehicle="Газель Next ВА5586-7",
        driver="Лабун Александр Чеславович",
        ttn_series="ЕМ",
        ttn_number="1926709",
        goods_accepted_by="Лабун Александр Чеславович",
        release_allowed_by="Шараев А.А.",
        shipper_handed_over_by="Шараев А.А.",
    )
    long_product_table = [
        ["Наименование товара", "Единица измерения", "Количество"],
        ["1", "2", "3"],
        *[[f"Товар {index}", "шт.", "1"] for index in range(1, 28)],
        ["Итого по странице", "X", "27"],
        ["Итого", "X", "99"],
    ]
    extracted = ExtractedTtnData(
        text="02 августа 2026 г.\nТОВАРНЫЙ РАЗДЕЛ",
        items=[],
        tables=[
            [["Грузоотправитель"]],
            [["Наименование товара"]],
            [["Операции"]],
            [["Расстояние"]],
            [["Товар 1"]],
            [["Товар 2"]],
            [["Товар 3"]],
            long_product_table,
        ],
    )

    write_ttn_excel(output, edit_data, extracted)

    workbook = load_workbook(output)
    sheet = workbook["Table 24"]
    assert sheet.max_row == len(long_product_table)
    assert sheet["A21"].value == "Товар 19"
    assert sheet["A21"]._style == sheet["A18"]._style
    assert sheet.row_dimensions[21].height == sheet.row_dimensions[18].height
    assert sheet["A30"].value == "Итого по странице"
    assert sheet["A30"]._style == sheet["A19"]._style
    assert sheet["A31"].value == "Итого"
    assert sheet["A31"]._style == sheet["A20"]._style


def test_write_ttn_excel_creates_extra_product_sheets(tmp_path: Path) -> None:
    output = tmp_path / "ttn_extra_product_pages.xlsx"
    edit_data = TtnEditData(
        vehicle="Газель Next ВА5586-7",
        driver="Лабун Александр Чеславович",
        ttn_series="ЕМ",
        ttn_number="1926709",
        goods_accepted_by="Лабун Александр Чеславович",
        release_allowed_by="Шараев А.А.",
        shipper_handed_over_by="Шараев А.А.",
    )
    product_table = [
        ["Наименование товара", "Единица измерения", "Количество"],
        ["1", "2", "3"],
        ["Товар", "шт.", "1"],
        ["Итого по странице", "X", "1"],
    ]
    extracted = ExtractedTtnData(
        text="02 августа 2026 г.\nТОВАРНЫЙ РАЗДЕЛ",
        items=[],
        tables=[
            [["Грузоотправитель"]],
            [["Наименование товара"]],
            [["Операции"]],
            [["Расстояние"]],
            product_table,
            product_table,
            product_table,
            product_table,
            product_table,
            product_table,
        ],
    )

    write_ttn_excel(output, edit_data, extracted)

    workbook = load_workbook(output)
    assert "Table 26" in workbook.sheetnames
    assert "Table 27" in workbook.sheetnames
    assert "Table 28" in workbook.sheetnames
    assert workbook["Table 26"]["A3"].value == "Товар"
    assert workbook["Table 26"]["A3"]._style == workbook["Table 24"]["A3"]._style
    assert workbook["Table 27"]["A1"].value == "Грузоотправитель"
    assert workbook["Table 27"]["A1"]._style == workbook["Table 25"]["A1"]._style
    assert workbook["Table 28"]["A3"].value == "Товар"
    assert workbook["Table 28"]["A3"]._style == workbook["Table 24"]["A3"]._style
