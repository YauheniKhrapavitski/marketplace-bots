from __future__ import annotations

import re
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ttn_bot.models import TtnEditData
from ttn_bot.pdf.extractor import ExtractedTtnData

DEFAULT_EXCEL_TEMPLATE_PATH = Path("config/excel_templates/ttn_ilovepdf_sample.xlsx")
PRODUCT_SHEET_START_NUMBER = 18
LAST_TEMPLATE_PRODUCT_SHEET = "Table 24"
LAST_TEMPLATE_SHIPPER_SHEET = "Table 25"

THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)


def write_ttn_excel(
    path: Path,
    edit_data: TtnEditData,
    extracted: ExtractedTtnData,
    template_path: Path | None = DEFAULT_EXCEL_TEMPLATE_PATH,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if template_path is not None and template_path.exists():
        _write_from_template(path, edit_data, extracted, template_path)
        return

    workbook = Workbook()
    workbook.remove(workbook.active)

    if extracted.tables:
        for index, table in enumerate(extracted.tables, start=1):
            sheet = workbook.create_sheet(f"Table {index}")
            _write_pdf_table(sheet, table)
    else:
        sheet = workbook.create_sheet("Table 1")
        _write_fallback_text(sheet, edit_data, extracted)

    workbook.save(path)


def validate_excel(path: Path) -> None:
    if not path.exists() or path.stat().st_size == 0:
        raise ValueError("Excel-файл не был создан")


def _write_from_template(
    path: Path,
    edit_data: TtnEditData,
    extracted: ExtractedTtnData,
    template_path: Path,
) -> None:
    workbook = load_workbook(template_path)
    blocks = _build_template_blocks(edit_data, extracted)
    _ensure_template_sheets(workbook, blocks)

    for sheet in workbook.worksheets:
        block = blocks.get(sheet.title)
        if block:
            _write_block(sheet, block)
        else:
            _clear_sheet_values(sheet)
        sheet.sheet_view.showGridLines = False

    workbook.save(path)


def _build_template_blocks(
    edit_data: TtnEditData,
    extracted: ExtractedTtnData,
) -> dict[str, list[list[str | int | float | None]]]:
    tables = extracted.tables
    first_page_tables = tables[:4]
    product_tables = tables[4:]
    text = extracted.text

    blocks: dict[str, list[list[str | int | float | None]]] = {
        "Table 1": _clean_table(first_page_tables[0]) if len(first_page_tables) > 0 else [],
        "Table 2": [["УНП"]],
        "Table 3": [[_vehicle_block(edit_data, text), _trailer_block()]],
        "Table 4": [[_driver_block(edit_data)]],
        "Table 5": [[_party_block(text)]],
        "Table 6": [[_basis_block(text)]],
        "Table 7": [[_loading_point_block(text), _unloading_point_block(text)]],
        "Table 8": [[_redirect_block()]],
        "Table 9": _clean_table(first_page_tables[1]) if len(first_page_tables) > 1 else [],
        "Table 10": [[_totals_block(text)]],
        "Table 11": [
            [_shipper_signature_block(edit_data, text), _carrier_signature_block(edit_data, text)]
        ],
        "Table 12": [
            ["Штамп", "№ пломбы\nII. ПОГРУЗОЧНО-РАЗГРУЗОЧНЫЕ ОПЕРАЦИИ", "Транспортные услуги"]
        ],
        "Table 13": _clean_table(first_page_tables[2]) if len(first_page_tables) > 2 else [],
        "Table 14": [[_other_info_block()]],
        "Table 15": _misc_left_block(first_page_tables[3]) if len(first_page_tables) > 3 else [],
        "Table 16": _misc_right_block(first_page_tables[3]) if len(first_page_tables) > 3 else [],
        "Table 17": [["Грузоотправитель"]],
        "Table 19": [["Грузоотправитель"]],
        "Table 21": [["Грузоотправитель"]],
        "Table 23": [["Грузоотправитель"]],
    }

    for index, table in enumerate(product_tables):
        blocks[_product_sheet_name(index)] = _clean_table(table)
        blocks[_shipper_sheet_name(index)] = [["Грузоотправитель"]]

    blocks.setdefault(LAST_TEMPLATE_SHIPPER_SHEET, [["Грузоотправитель"]])

    return blocks


def _ensure_template_sheets(
    workbook: Workbook,
    blocks: dict[str, list[list[str | int | float | None]]],
) -> None:
    for sheet_name in sorted(blocks, key=_table_number):
        if sheet_name in workbook.sheetnames:
            continue
        source_sheet_name = (
            LAST_TEMPLATE_PRODUCT_SHEET
            if _is_product_sheet(sheet_name)
            else LAST_TEMPLATE_SHIPPER_SHEET
        )
        source_sheet = workbook[source_sheet_name]
        copied_sheet = workbook.copy_worksheet(source_sheet)
        copied_sheet.title = sheet_name


def _product_sheet_name(index: int) -> str:
    return f"Table {PRODUCT_SHEET_START_NUMBER + index * 2}"


def _shipper_sheet_name(index: int) -> str:
    return f"Table {PRODUCT_SHEET_START_NUMBER - 1 + index * 2}"


def _is_product_sheet(sheet_name: str) -> bool:
    number = _table_number(sheet_name)
    return number >= PRODUCT_SHEET_START_NUMBER and number % 2 == 0


def _table_number(sheet_name: str) -> int:
    match = re.search(r"\d+$", sheet_name)
    return int(match.group(0)) if match else 0


def _clear_sheet_values(sheet: Worksheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None


def _write_block(sheet: Worksheet, rows: list[list[str | int | float | None]]) -> None:
    if _is_product_sheet(sheet.title):
        _extend_product_table_styles(sheet, rows)

    _clear_sheet_values(sheet)

    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index)
            if isinstance(cell, MergedCell):
                continue
            cell.value = value


def _extend_product_table_styles(
    sheet: Worksheet,
    rows: list[list[str | int | float | None]],
) -> None:
    template_max_row = sheet.max_row
    template_max_column = sheet.max_column
    product_style_row = _product_style_row(sheet)
    page_total_style_row = _style_row_by_label(sheet, "Итого по странице") or template_max_row
    total_style_row = _style_row_by_label(sheet, "Итого") or page_total_style_row

    for row_index, row in enumerate(rows, start=1):
        source_row = _source_style_row_for_product_row(
            row,
            row_index,
            product_style_row,
            page_total_style_row,
            total_style_row,
        )
        if source_row is None:
            continue
        _copy_row_format(sheet, source_row, row_index, template_max_column)


def _source_style_row_for_product_row(
    row: list[str | int | float | None],
    row_index: int,
    product_style_row: int,
    page_total_style_row: int,
    total_style_row: int,
) -> int | None:
    if row_index <= 2:
        return None
    label = str(row[0] or "").strip()
    if label == "Итого":
        return total_style_row
    if label.startswith("Итого"):
        return page_total_style_row
    return product_style_row


def _product_style_row(sheet: Worksheet) -> int:
    for row_index in range(sheet.max_row, 2, -1):
        value = str(sheet.cell(row_index, 1).value or "").strip()
        if value and not value.startswith("Итого"):
            return row_index
    return min(sheet.max_row, 3)


def _style_row_by_label(sheet: Worksheet, label: str) -> int | None:
    for row_index in range(1, sheet.max_row + 1):
        value = str(sheet.cell(row_index, 1).value or "").strip()
        if value == label:
            return row_index
    return None


def _copy_row_format(
    sheet: Worksheet,
    source_row: int,
    target_row: int,
    max_column: int,
) -> None:
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height
    for column_index in range(1, max_column + 1):
        source_cell = sheet.cell(source_row, column_index)
        target_cell = sheet.cell(target_row, column_index)
        if isinstance(target_cell, MergedCell):
            continue
        target_cell._style = copy(source_cell._style)
        if source_cell.has_style:
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.number_format = source_cell.number_format
            target_cell.protection = copy(source_cell.protection)


def _clean_table(table: list[list[str | None]]) -> list[list[str | int | float | None]]:
    return [[_clean_excel_value(cell) for cell in row] for row in table]


def _vehicle_block(edit_data: TtnEditData, text: str) -> str:
    return (
        f"{_document_date(text)}\n"
        f"Автомобиль {edit_data.vehicle}\n"
        "(марка, государственный номер)"
    )


def _trailer_block() -> str:
    return "Прицеп \n(марка, государственный номер)"


def _driver_block(edit_data: TtnEditData) -> str:
    return (
        "К Путевому листу № \n"
        f"Водитель {edit_data.driver}\n"
        "(фамилия и инициалы)"
    )


def _party_block(text: str) -> str:
    return _lines_between(
        text,
        "Заказчик автомобильной перевозки",
        "Основание отпуска",
        fallback="Грузоотправитель\nГрузополучатель",
    )


def _basis_block(text: str) -> str:
    return _lines_between(
        text,
        "Основание отпуска",
        "Пункт погрузки",
        fallback="Основание отпуска",
    )


def _loading_point_block(text: str) -> str:
    return _lines_between(
        text,
        "Пункт погрузки",
        "Пункт разгрузки",
        fallback="Пункт погрузки\n(адрес)",
    )


def _unloading_point_block(text: str) -> str:
    return _lines_between(
        text,
        "Пункт разгрузки",
        "Переадресовка",
        fallback="Пункт разгрузки\n(адрес)",
    )


def _redirect_block() -> str:
    return (
        "Переадресовка \n"
        "(новый пункт разгрузки, фамилия, инициалы, подпись лица, "
        "принявшего решение о переадресовке)\n"
        "I. ТОВАРНЫЙ РАЗДЕЛ"
    )


def _totals_block(text: str) -> str:
    lines = [
        line
        for line in text.splitlines()
        if line.startswith("Всего сумма НДС") or line.startswith("Всего стоимость")
    ]
    return "\n".join(lines) if lines else "Всего сумма НДС\nВсего стоимость c НДС"


def _shipper_signature_block(edit_data: TtnEditData, text: str) -> str:
    mass_line = _first_line_starting(text, "Всего масса груза") or "Всего масса груза"
    return (
        f"{mass_line}\n(прописью)\n"
        f"Отпуск разрешил {edit_data.release_allowed_by}\n"
        "(должность,\nфамилия, инициалы, подпись)\n"
        f"Сдал грузоотправитель {edit_data.shipper_handed_over_by}\n"
        "(должность,\nфамилия, инициалы, подпись)\n"
        "№ пломбы"
    )


def _carrier_signature_block(edit_data: TtnEditData, text: str) -> str:
    places_line = (
        _first_line_starting(text, "Всего количество грузовых мест")
        or "Всего количество грузовых мест"
    )
    return (
        f"{places_line}\n(прописью)\n"
        f"Товар к перевозке принял {edit_data.goods_accepted_by}\n"
        "(должность,\nфамилия, инициалы, подпись)\n"
        "по доверенности \n(номер, дата)\n"
        "выданной \n(наименование организации)\n"
        "Принял грузополучатель \n(должность,"
    )


def _other_info_block() -> str:
    return (
        "III. ПРОЧИЕ СВЕДЕНИЯ (заполняются перевозчиком)\n"
        "Отметки о составленных актах:\n"
        "Таксировка\n"
        "С товаром переданы документы: товарно-транспортная накладная и приложение к ней."
    )


def _misc_left_block(table: list[list[str | None]]) -> list[list[str | int | float | None]]:
    cleaned = _clean_table(table)
    return [row[:11] for row in cleaned[:5]]


def _misc_right_block(table: list[list[str | None]]) -> list[list[str | int | float | None]]:
    cleaned = _clean_table(table)
    return [row[10:22] for row in cleaned[5:12]]


def _document_date(text: str) -> str:
    match = re.search(r"\b\d{1,2}\s+[А-Яа-яёЁ]+\s+\d{4}\s+г\.", text)
    return match.group(0) if match else ""


def _first_line_starting(text: str, prefix: str) -> str | None:
    return next(
        (line.strip() for line in text.splitlines() if line.strip().startswith(prefix)),
        None,
    )


def _lines_between(text: str, start: str, end: str, fallback: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    start_index = next((index for index, line in enumerate(lines) if line.startswith(start)), None)
    if start_index is None:
        return fallback
    end_index = next(
        (
            index
            for index, line in enumerate(lines[start_index + 1 :], start=start_index + 1)
            if line.startswith(end)
        ),
        len(lines),
    )
    return "\n".join(lines[start_index:end_index]) or fallback


def _write_pdf_table(sheet: Worksheet, table: list[list[str | None]]) -> None:
    max_cols = max((len(row) for row in table), default=1)
    for row_index, row in enumerate(table, start=1):
        for column_index in range(1, max_cols + 1):
            value = row[column_index - 1] if column_index <= len(row) else None
            cell = sheet.cell(row=row_index, column=column_index, value=_clean_excel_value(value))
            cell.font = Font(name="Times New Roman", size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = THIN_BORDER

    _style_header_rows(sheet)
    _autosize_like_pdf_converter(sheet)
    sheet.sheet_view.showGridLines = False


def _write_fallback_text(
    sheet: Worksheet,
    edit_data: TtnEditData,
    extracted: ExtractedTtnData,
) -> None:
    lines = [
        f"Автомобиль: {edit_data.vehicle}",
        f"Водитель: {edit_data.driver}",
        f"Серия и номер ТТН: {edit_data.series_and_number}",
        f"Отпуск разрешил: {edit_data.release_allowed_by}",
        "",
        extracted.text,
    ]
    sheet["A1"] = "\n".join(lines)
    sheet["A1"].font = Font(name="Times New Roman", size=9)
    sheet["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 132.44
    sheet.row_dimensions[1].height = 120


def _style_header_rows(sheet: Worksheet) -> None:
    for row in range(1, min(sheet.max_row, 2) + 1):
        for cell in sheet[row]:
            cell.font = Font(name="Times New Roman", size=9, bold=True)


def _autosize_like_pdf_converter(sheet: Worksheet) -> None:
    max_column = sheet.max_column
    if max_column == 1:
        sheet.column_dimensions["A"].width = 132.44
        return

    for column_index in range(1, max_column + 1):
        values = [
            str(sheet.cell(row=row, column=column_index).value or "")
            for row in range(1, sheet.max_row + 1)
        ]
        longest = max((len(part) for value in values for part in value.splitlines()), default=0)
        if column_index == 1 and max_column >= 10:
            width = 45.78
        else:
            width = min(max(longest * 0.9 + 2, 5.33), 22.0)
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index in range(1, sheet.max_row + 1):
        max_lines = max(
            len(str(sheet.cell(row=row_index, column=column).value or "").splitlines())
            for column in range(1, max_column + 1)
        )
        sheet.row_dimensions[row_index].height = max(12, min(14 + (max_lines - 1) * 8, 60))


def _clean_excel_value(value: str | None) -> str | int | float | None:
    if value is None:
        return None
    cleaned = value.replace("\r", "\n").strip()
    if cleaned == "":
        return None
    if "\n" in cleaned:
        return cleaned
    if cleaned.isdigit():
        return int(cleaned)
    normalized = cleaned.replace(" ", "")
    if normalized.count(".") == 1 and normalized.replace(".", "").isdigit():
        return float(normalized)
    return cleaned
